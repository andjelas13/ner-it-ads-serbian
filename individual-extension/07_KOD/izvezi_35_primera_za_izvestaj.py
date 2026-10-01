#!/usr/bin/env python3
"""Izvezi tačno 35 slučajeva prikazanih u kvalitativnom delu izveštaja.
ž
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "06_RUCNA_PROVERA" / "slucajevi_1114.json"
REVIEW_TABLE = ROOT / "06_RUCNA_PROVERA" / "pregled_uzorka_100.xlsx"
OUTPUT = ROOT / "04_REZULTATI" / "izabrani_primeri_35_za_izvestaj.csv"

DECISION_CODES = {
    "Bolja je ručna anotacija": "RUCNA",
    "Bolja je LLM anotacija": "LLM",
    "Obe su prihvatljive": "OBE",
    "Nijedna nije potpuno ispravna": "NIJEDNA",
}


def read_report_decisions() -> dict[str, tuple[str, str]]:
    """Vraća konačnu odluku i obrazloženje za slučajeve označene za izveštaj.

    """
    sheet = load_workbook(REVIEW_TABLE, read_only=True)["Uzorak 100"]
    rows = list(sheet.iter_rows(values_only=True))
   
    header = list(rows[0])
    case_i = header.index("Slučaj")
    report_i = header.index("U izveštaju")
    final_i = header.index("Konačna odluka")
    reason_i = header.index("Obrazloženje")

    decisions: dict[str, tuple[str, str]] = {}
    for values in rows[1:]:
      
        if not values or values[report_i] != "DA":
            continue
       
        if values[final_i] not in DECISION_CODES:
            raise ValueError(f"Slučaj {values[case_i]} nema upisanu konačnu odluku.")
        # Obrazloženje nije obavezno; ako ga u tabeli nema, kolona u izvozu ostaje prazna.
        decisions[values[case_i]] = (DECISION_CODES[values[final_i]], values[reason_i] or "")
   
    if not decisions:
        raise SystemExit(
            f"U tabeli {REVIEW_TABLE.name} još nijedan slučaj nije označen za izveštaj "
            "(kolona „U izveštaju“ = DA)."
        )
    return decisions


def main() -> int:
    
    decisions = read_report_decisions()
    payload = json.loads(SOURCE.read_text(encoding="utf-8-sig"))
    selected = [case for case in payload["cases"] if case["case_id"] in decisions]
   
    if len(selected) != 35 or {case["case_id"] for case in selected} != set(decisions):
        raise ValueError("Nije pronađeno svih 35 izabranih slučajeva")

    selected.sort(key=lambda case: case["case_id"])

    columns = [
        "case_id", "ad_id", "field", "structural_category", "primary_label",
        "context", "human_spans", "llm_spans", "decision_for_report",
        "explanation_for_report",
    ]
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
  
    with OUTPUT.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for case in selected:
           
            decision, reason = decisions[case["case_id"]]
            writer.writerow(
                {
                    "case_id": case["case_id"],
                    "ad_id": case["ad_id"],
                    "field": case["field"],
                    "structural_category": case["structural_category"],
                    "primary_label": case["primary_label"],
                    "context": case["context"],
                   
                    "human_spans": json.dumps(case["human_spans"], ensure_ascii=False),
                    "llm_spans": json.dumps(case["llm_spans"], ensure_ascii=False),
                    "decision_for_report": decision,
                    "explanation_for_report": reason,
                }
            )
   
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
