#!/usr/bin/env python3
"""Formira reproduktivan stratifikovani uzorak od 100 slučajeva i tabelu za pregled.


"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import re
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_INPUT = ROOT / "06_RUCNA_PROVERA" / "slucajevi_1114.json"
DEFAULT_OUTPUT_JSON = ROOT / "06_RUCNA_PROVERA" / "uzorak_100.json"
DEFAULT_OUTPUT_CSV = ROOT / "06_RUCNA_PROVERA" / "uzorak_100.csv"
DEFAULT_TABLE = ROOT / "06_RUCNA_PROVERA" / "pregled_uzorka_100.xlsx"

SEED = 20260914

# Uzorak je stratifikovan: kvota unapred određuje koliko slučajeva ulazi iz
# svake strukturne kategorije razlike. Kvote nisu srazmerne veličini
# kategorija. Retke kategorije (POGRESNA_OZNAKA 17, GRANICA_I_OZNAKA 5,
# GRANICA_POMERENA 2, SLOZENO 1) uzimaju se u celosti, da nijedna vrsta
# razlike ne ostane nepregledana, a velike kategorije se ograničavaju kako ne
# bi popunile ceo uzorak. Zbir kvota je tačno 100.
QUOTAS = {
    "POGRESNA_OZNAKA": 17,
    "GRANICA_I_OZNAKA": 5,
    "GRANICA_POMERENA": 2,
    "SLOZENO": 1,
    "SUVISNO": 18,
    "PROPUSTENO": 15,
    "GRANICA_SIRA": 15,
    "GRANICA_UZA": 8,
    "FRAGMENTACIJA": 10,
    "SPAJANJE": 9,
}

EXPECTED_SELECTION_SHA256 = (
    "b8e22edd1fd56dfebe0531bfe22502093473d4a7aa1c41cd74c41dccb7d29856"
)

CASE_COLUMNS = [
    "Redni broj u uzorku",
    "Slučaj",
    "Oglas",
    "Polje",
    "Strukturna kategorija",
    "Oznaka",
    "Kontekst oglasa",
    "Ručna anotacija",
    "LLM anotacija",
]
REVIEW_COLUMNS = [
    "Odluka u prvom prolazu",
    "Beleška iz prvog prolaza",
    "U izveštaju",
    "Konačna odluka",
    "Obrazloženje",
]

DECISIONS = {
    "RUCNA": "Bolja je ručna anotacija",
    "LLM": "Bolja je LLM anotacija",
    "OBE": "Obe su prihvatljive",
    "NIJEDNA": "Nijedna nije potpuno ispravna",
}

CATEGORY_NAMES = {
    "SUVISNO": "SUVIŠNO",
    "PROPUSTENO": "PROPUŠTENO",
    "GRANICA_SIRA": "GRANICA ŠIRA",
    "GRANICA_UZA": "GRANICA UŽA",
    "GRANICA_POMERENA": "GRANICA POMERENA",
    "POGRESNA_OZNAKA": "POGREŠNA OZNAKA",
    "GRANICA_I_OZNAKA": "GRANICA I OZNAKA",
    "FRAGMENTACIJA": "FRAGMENTACIJA",
    "SPAJANJE": "SPAJANJE",
    "SLOZENO": "SLOŽENO",
}

HEADER_FILL = "7A1F2E"


def balanced_pick(
    rows: list[dict], quota: int, rng: random.Random
) -> list[dict]:
    """Bira slučajeve uz ravnotežu oglasa, oznaka i delova skupa.

    

    Zato se u svakom koraku bira slučaj koji je najmanje sličan već izabranima,
    po tri merila redom:
    1. slučaj iz oglasa koji još nije zastupljen ima prednost (kazna 2 ako je
       oglas već viđen, 0 ako nije);
    2. zatim slučaj sa oznakom (CPU, GPU, RAM, SKLADIŠTE, CENA) koja je do tada
       najređe izabrana;
    3. zatim slučaj iz dela skupa (1–4) koji je najređe izabran.

    Ako je i posle toga više slučajeva jednako dobro, uzima se prvi po redu u
    listi koja je unapred izmešana zadatim semenom. Tako je odluka nasumična,
    ali uvek ista pri svakom pokretanju.
    """
    pool = list(rows)
    rng.shuffle(pool)
    chosen: list[dict] = []
    label_counts: Counter[str] = Counter()
    part_counts: Counter[int] = Counter()
    seen_ads: set[int] = set()

    while pool and len(chosen) < quota:
        # Manja vrednost ključa znači bolji kandidat, pa ``min`` bira slučaj
    
        best_index = min(
            range(len(pool)),
            key=lambda i: (
                2 if pool[i]["ad_id"] in seen_ads else 0,
                label_counts[pool[i]["primary_label"]],
                part_counts[pool[i]["part"]],
                i,
            ),
        )
        # Izabrani slučaj se vadi iz liste kandidata da ne bi bio izabran

        row = pool.pop(best_index)
        chosen.append(row)
        seen_ads.add(row["ad_id"])
        label_counts[row["primary_label"]] += 1
        part_counts[row["part"]] += 1

    # Do ovoga dolazi samo ako kategorija ima manje slučajeva nego što traži

    if len(chosen) != quota:
        raise ValueError(
            f"Za kategoriju je traženo {quota} slučajeva, a dostupno je {len(rows)}."
        )
    return chosen


def selection_sha256(cases: list[dict]) -> str:
    
    case_ids = sorted(case["case_id"] for case in cases)
    return hashlib.sha256("\n".join(case_ids).encode("utf-8")).hexdigest()


def relative(path: Path) -> str:
 
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return path.as_posix()


def matches_saved_sample(selected: list[dict], saved_path: Path) -> bool | None:
  
    if not saved_path.exists():
        return None
    saved = json.loads(saved_path.read_text(encoding="utf-8-sig"))["cases"]
    return sorted(case["case_id"] for case in saved) == sorted(case["case_id"] for case in selected)


def form_sample(input_path: Path) -> tuple[list[dict], list[dict], str]:
    """Formira uzorak od 100 slučajeva iz baze svih razlika.

    Vraća tri stvari: izabranih 100 slučajeva (svaki sa dodatim rednim brojem
    u uzorku).


    """
    payload = json.loads(input_path.read_text(encoding="utf-8-sig"))
    all_cases = payload["cases"]

    if len({case["case_id"] for case in all_cases}) != len(all_cases):
        raise ValueError("Ulaz sadrži ponovljene case_id vrednosti.")

    ordered_cases = sorted(all_cases, key=lambda case: case["case_id"])
    rng = random.Random(SEED)
    selected: list[dict] = []

    for category, quota in QUOTAS.items():
        candidates = [
            case
            for case in ordered_cases
            if case["structural_category"] == category
        ]
        selected.extend(balanced_pick(candidates, quota, rng))

    
    selected.sort(key=lambda case: case["case_id"])
    selected = [dict(case) for case in selected]
    for order, case in enumerate(selected, start=1):
        case["sample_order"] = order

    if len(selected) != 100 or len({case["case_id"] for case in selected}) != 100:
        raise RuntimeError("Uzorak ne sadrži tačno 100 jedinstvenih slučajeva.")

    digest = selection_sha256(selected)
    if digest != EXPECTED_SELECTION_SHA256:
        raise RuntimeError(
            "Dobijeni uzorak se ne podudara sa sačuvanim uzorkom. "
            f"Očekivani SHA-256: {EXPECTED_SELECTION_SHA256}; dobijeni: {digest}."
        )

    source_digest = hashlib.sha256(input_path.read_bytes()).hexdigest()
    return selected, all_cases, source_digest


def write_outputs(
    selected: list[dict],
    population_size: int,
    source_digest: str,
    output_json: Path,
    output_csv: Path,
) -> None:
    
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_csv.parent.mkdir(parents=True, exist_ok=True)

    metadata = {
        "description": "Stratifikovani uzorak za kvalitativnu proveru razlika",
        "seed": SEED,
        "population_size": population_size,
        "sample_size": len(selected),
        "quotas": QUOTAS,
        "selection_algorithm": (
            "U svakoj strukturnoj kategoriji prednost imaju novi oglasi, "
            "zatim slabije zastupljene oznake i delovi skupa."
        ),
        "selected_case_ids_sha256": selection_sha256(selected),
        "source_file_sha256": source_digest,
        "review_table": relative(DEFAULT_TABLE),
    }
    output_json.write_text(
        json.dumps(
            {"metadata": metadata, "cases": selected},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
        newline="\n",
    )

    fields = [
        "sample_order",
        "case_id",
        "ad_id",
        "field",
        "part",
        "structural_category",
        "primary_label",
        "context",
    ]
    with output_csv.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(selected)


def format_spans(spans: list[dict]) -> str:
   
    if not spans:
        return "—"
    parts = []
    for span in spans:
        text = re.sub(r"\s+", " ", span["text"])
        label = span["label"].replace("SKLADISTE", "SKLADIŠTE")
        parts.append(f"⟦{text}⟧ {label}")
    return "; ".join(parts)


def case_row(case: dict, order: int | str) -> list:
  
    return [
        order,
        case["case_id"],
        int(case["ad_id"]),
        "naslov" if case["field"] == "title" else "telo",
        CATEGORY_NAMES[case["structural_category"]],
        case["primary_label"].replace("SKLADISTE", "SKLADIŠTE"),
        re.sub(r"\s+", " ", case["context"]).strip(),
        format_spans(case["human_spans"]),
        format_spans(case["llm_spans"]),
    ]


def write_review_table(selected: list[dict], table_path: Path) -> None:
   
   
    
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Uzorak 100"
    sheet.append(CASE_COLUMNS + REVIEW_COLUMNS)
    # Redovi idu po rednom broju u uzorku, a kolone za odluke ostaju prazne.
    for case in sorted(selected, key=lambda item: item["sample_order"]):
        sheet.append(case_row(case, case["sample_order"]) + [""] * len(REVIEW_COLUMNS))

    widths = [9, 8, 7, 7, 17, 11, 45, 28, 28, 18, 24, 9, 18, 70]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
          
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
  
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions

    last_row = len(selected) + 50
    decision_list = DataValidation(type="list", formula1='"' + ",".join(DECISIONS.values()) + '"', allow_blank=True)
    report_list = DataValidation(type="list", formula1='"DA,NE"', allow_blank=True)
    sheet.add_data_validation(decision_list)
    sheet.add_data_validation(report_list)
    for name in ("Odluka u prvom prolazu", "Konačna odluka"):
        letter = get_column_letter((CASE_COLUMNS + REVIEW_COLUMNS).index(name) + 1)
        decision_list.add(f"{letter}2:{letter}{last_row}")
    letter = get_column_letter((CASE_COLUMNS + REVIEW_COLUMNS).index("U izveštaju") + 1)
    report_list.add(f"{letter}2:{letter}{last_row}")

    table_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(table_path)


def check_review_table(selected: list[dict], all_cases: list[dict], table_path: Path) -> dict:
   
    from openpyxl import load_workbook

   
    sheet = load_workbook(table_path, read_only=True)["Uzorak 100"]
    rows = list(sheet.iter_rows(values_only=True))
    if list(rows[0][: len(CASE_COLUMNS + REVIEW_COLUMNS)]) != CASE_COLUMNS + REVIEW_COLUMNS:
        raise RuntimeError(f"Zaglavlje tabele {relative(table_path)} nije očekivano.")

    by_id = {case["case_id"]: case for case in all_cases}
    in_sample = {case["case_id"]: case for case in selected}
    seen: Counter[str] = Counter()
    added = 0
    for values in rows[1:]:
        
        if not any(values):
            continue
        case_id = values[1]
        if case_id not in by_id:
            raise RuntimeError(f"Slučaj {case_id} iz tabele ne postoji u bazi razlika.")
        seen[case_id] += 1
        if case_id in in_sample:
            expected = case_row(in_sample[case_id], in_sample[case_id]["sample_order"])
        else:
            expected = case_row(by_id[case_id], "—")
            added += 1
        found = [value if value is not None else "" for value in values[: len(CASE_COLUMNS)]]
        if [str(value) for value in found] != [str(value) for value in expected]:
            raise RuntimeError(f"Podaci o slučaju {case_id} u tabeli ne odgovaraju uzorku.")
       
        for name in ("Odluka u prvom prolazu", "Konačna odluka"):
            value = values[(CASE_COLUMNS + REVIEW_COLUMNS).index(name)]
            if value and value not in DECISIONS.values():
                raise RuntimeError(f"Nepoznata odluka za {case_id}: {value}")

    duplicates = [case_id for case_id, count in seen.items() if count > 1]
    missing = sorted(set(in_sample) - set(seen))
    if duplicates or missing:
        raise RuntimeError(f"Tabela ima ponovljene ({duplicates}) ili nedostajuće ({missing}) slučajeve.")

    final_index = (CASE_COLUMNS + REVIEW_COLUMNS).index("Konačna odluka")
    report_index = (CASE_COLUMNS + REVIEW_COLUMNS).index("U izveštaju")
    return {
        "rows": sum(seen.values()),
        "added_from_full_base": added,
        "final_decisions": sum(1 for values in rows[1:] if values and values[final_index]),
        "in_report": sum(1 for values in rows[1:] if values and values[report_index] == "DA"),
    }


def parse_args() -> argparse.Namespace:
   
    parser = argparse.ArgumentParser(
        description="Formira isti stratifikovani uzorak od 100 slučajeva i tabelu za pregled."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="baza svih slučajeva razlike")
    parser.add_argument("--output-json", type=Path, default=DEFAULT_OUTPUT_JSON, help="izlazni JSON uzorka")
    parser.add_argument("--output-csv", type=Path, default=DEFAULT_OUTPUT_CSV, help="izlazni CSV uzorka")
    parser.add_argument("--tabela", type=Path, default=DEFAULT_TABLE, help="tabela za pregled (pravi se ako ne postoji)")
    return parser.parse_args()


def main() -> int:
  
    args = parse_args()
    selected, all_cases, source_digest = form_sample(args.input)

    if matches_saved_sample(selected, args.output_json) is False:
        raise RuntimeError(f"Izbor se ne podudara sa uzorkom u {relative(args.output_json)}.")
    write_outputs(
        selected,
        len(all_cases),
        source_digest,
        args.output_json,
        args.output_csv,
    )

    if args.tabela.exists():
        table_status = {"status": "postojeća tabela je proverena i nije menjana"}
        table_status.update(check_review_table(selected, all_cases, args.tabela))
    else:
        write_review_table(selected, args.tabela)
        table_status = {"status": "napravljena je nova tabela sa praznim kolonama za odluke"}

    print(
        json.dumps(
            {
                "input": relative(args.input),
                "output_json": relative(args.output_json),
                "output_csv": relative(args.output_csv),
                "review_table": relative(args.tabela),
                "review_table_check": table_status,
                "sample_size": len(selected),
                "by_category": dict(
                    Counter(case["structural_category"] for case in selected)
                ),
                "selected_case_ids_sha256": selection_sha256(selected),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
